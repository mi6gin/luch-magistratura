"""Render a forecasting report document into Ray OS PDF and PPTX artifacts."""

from __future__ import annotations

import importlib.util
import os
import re
import tempfile
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping, Sequence
from xml.sax.saxutils import escape

if __package__:
    from .forecasting import build_report_document
else:
    from forecasting import build_report_document


NAVY = "#0B1020"
CHARCOAL = "#171D2F"
CORAL = "#FF6B5E"
LIME = "#C8F169"
POWDER = "#A8D8FF"
WHITE = "#F6F7FB"
MUTED = "#98A3B8"
GRID = "#31394F"


def _normalize_formats(formats: Sequence[str] | str | None) -> list[str]:
    if formats is None:
        values = ["pdf"]
    elif isinstance(formats, str):
        values = [value.strip().lower() for value in formats.split(",") if value.strip()]
    else:
        values = [str(value).strip().lower() for value in formats if str(value).strip()]
    values = list(dict.fromkeys(values))
    if not values:
        raise ValueError("At least one report format is required.")
    invalid = sorted(set(values) - {"pdf", "pptx"})
    if invalid:
        raise ValueError(f"Unsupported report formats: {', '.join(invalid)}")
    return values


def _check_dependencies(formats: Sequence[str]) -> None:
    required = {"matplotlib"}
    if "pdf" in formats:
        required.add("reportlab")
    if "pptx" in formats:
        required.add("pptx")
    missing = sorted(module for module in required if importlib.util.find_spec(module) is None)
    if missing:
        package_names = ["python-pptx" if module == "pptx" else module for module in missing]
        raise RuntimeError(f"Missing report dependencies: {', '.join(package_names)}")


def _output_directory() -> Path:
    configured = os.getenv("ML_REPORTS_PATH")
    path = (
        Path(configured).expanduser()
        if configured
        else Path(__file__).resolve().parent.parent / "storage" / "app" / "reports"
    )
    path = path.resolve()
    path.mkdir(parents=True, exist_ok=True)
    return path


def _format_number(value: Any) -> str:
    return f"{float(value):,.0f}".replace(",", " ")


def _format_money(value: Any) -> str:
    return f"{_format_number(value)} KZT"


def _short(value: Any, length: int = 28) -> str:
    text = str(value)
    return text if len(text) <= length else f"{text[: length - 1]}..."


def _create_charts(document: Mapping[str, Any], work_dir: Path) -> dict[str, Path]:
    import matplotlib

    matplotlib.use("Agg", force=True)
    import matplotlib.pyplot as plt
    import numpy as np
    import pandas as pd

    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "DejaVu Sans"],
            "axes.facecolor": CHARCOAL,
            "figure.facecolor": NAVY,
            "text.color": WHITE,
            "axes.labelcolor": MUTED,
            "xtick.color": MUTED,
            "ytick.color": MUTED,
            "axes.edgecolor": GRID,
        }
    )

    actions = document["product_actions"]
    dates = pd.to_datetime(actions[0]["forecast"]["dates"])
    q10 = np.sum([item["forecast"]["q10"] for item in actions], axis=0)
    q50 = np.sum([item["forecast"]["q50"] for item in actions], axis=0)
    q90 = np.sum([item["forecast"]["q90"] for item in actions], axis=0)

    portfolio_path = work_dir / f"portfolio-{uuid.uuid4().hex}.png"
    fig, axis = plt.subplots(figsize=(12, 4.6))
    axis.fill_between(dates, q10, q90, color=POWDER, alpha=0.16, linewidth=0)
    axis.plot(dates, q50, color=CORAL, linewidth=3.2, label="Спрос q50")
    axis.plot(dates, q90, color=POWDER, linewidth=1.2, alpha=0.8, label="q90")
    axis.set_title("Сигнал спроса на 30 дней", loc="left", fontsize=17, fontweight="bold")
    axis.grid(axis="y", color=GRID, alpha=0.7, linewidth=0.7)
    axis.spines[["top", "right"]].set_visible(False)
    axis.legend(frameon=False, loc="upper right")
    fig.autofmt_xdate(rotation=0)
    fig.tight_layout()
    fig.savefig(portfolio_path, dpi=170, bbox_inches="tight", facecolor=NAVY)
    plt.close(fig)

    categories = document["category_summaries"]
    category_path = work_dir / f"category-{uuid.uuid4().hex}.png"
    labels = [item["category"] for item in categories]
    medians = np.array([item["forecast_30_q50"] for item in categories], dtype=float)
    uppers = np.array([item["forecast_30_q90"] for item in categories], dtype=float)
    positions = np.arange(len(labels))
    fig, axis = plt.subplots(figsize=(10.5, 4.7))
    axis.barh(positions, uppers, color=GRID, height=0.58, label="q90")
    axis.barh(positions, medians, color=LIME, height=0.58, label="q50")
    axis.set_yticks(positions, labels=labels)
    axis.invert_yaxis()
    axis.set_title("Спрос по категориям", loc="left", fontsize=17, fontweight="bold")
    axis.grid(axis="x", color=GRID, alpha=0.7, linewidth=0.7)
    axis.spines[["top", "right", "left"]].set_visible(False)
    axis.legend(frameon=False, loc="lower right")
    fig.tight_layout()
    fig.savefig(category_path, dpi=170, bbox_inches="tight", facecolor=NAVY)
    plt.close(fig)

    scatter_path = work_dir / f"portfolio-map-{uuid.uuid4().hex}.png"
    cover = np.array([min(float(item["days_cover"]), 90.0) for item in actions])
    demand = np.array([float(item["forecast_30_q50"]) for item in actions])
    orders = np.array([float(item["order_quantity"]) for item in actions])
    sizes = 45.0 + 260.0 * orders / max(float(orders.max()), 1.0)
    colors = [CORAL if item["order_quantity"] > 0 else POWDER for item in actions]
    fig, axis = plt.subplots(figsize=(10.5, 4.7))
    axis.scatter(cover, demand, s=sizes, c=colors, alpha=0.82, edgecolors=NAVY, linewidths=0.8)
    for item, x_value, y_value in zip(actions[:8], cover[:8], demand[:8]):
        axis.annotate(str(item["product_id"]), (x_value, y_value), color=WHITE, fontsize=8, ha="center", va="center")
    axis.axvline(14, color=LIME, linestyle="--", linewidth=1.2, alpha=0.8)
    axis.set_xlabel("Дни покрытия (до 90)")
    axis.set_ylabel("Спрос q50 на 30 дней")
    axis.set_title("Карта товарного портфеля", loc="left", fontsize=17, fontweight="bold")
    axis.grid(color=GRID, alpha=0.6, linewidth=0.7)
    axis.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(scatter_path, dpi=170, bbox_inches="tight", facecolor=NAVY)
    plt.close(fig)
    return {"portfolio": portfolio_path, "category": category_path, "scatter": scatter_path}


def _register_pdf_fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    regular_name = "RayOS-Regular"
    bold_name = "RayOS-Bold"
    if regular_name in pdfmetrics.getRegisteredFontNames() and bold_name in pdfmetrics.getRegisteredFontNames():
        return regular_name, bold_name

    candidates = [
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
        (
            Path("/usr/local/share/fonts/DejaVuSans.ttf"),
            Path("/usr/local/share/fonts/DejaVuSans-Bold.ttf"),
        ),
    ]
    selected: tuple[Path, Path] | None = next(
        ((regular, bold) for regular, bold in candidates if regular.is_file() and bold.is_file()),
        None,
    )
    if selected is None:
        from matplotlib import font_manager

        regular = Path(font_manager.findfont("DejaVu Sans", fallback_to_default=True))
        bold = Path(
            font_manager.findfont(
                font_manager.FontProperties(family="DejaVu Sans", weight="bold"),
                fallback_to_default=True,
            )
        )
        selected = (regular, bold)
    pdfmetrics.registerFont(TTFont(regular_name, str(selected[0])))
    pdfmetrics.registerFont(TTFont(bold_name, str(selected[1])))
    pdfmetrics.registerFontFamily(
        "RayOS",
        normal=regular_name,
        bold=bold_name,
        italic=regular_name,
        boldItalic=bold_name,
    )
    return regular_name, bold_name


def _render_pdf(document: Mapping[str, Any], destination: Path, charts: Mapping[str, Path]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Image,
        LongTable,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    regular, bold = _register_pdf_fonts()
    page_size = (13.333 * inch, 7.5 * inch)
    metadata = document["metadata"]
    warning_count = len(document["warnings"])

    def color(value: str):
        return colors.HexColor(value)

    def on_page(canvas, doc) -> None:
        width, height = page_size
        canvas.saveState()
        canvas.setFillColor(color(NAVY))
        canvas.rect(0, 0, width, height, fill=1, stroke=0)
        canvas.setStrokeColor(color(GRID))
        canvas.line(30, height - 33, width - 30, height - 33)
        canvas.line(30, 25, width - 30, 25)
        canvas.setFillColor(color(POWDER))
        canvas.setFont(bold, 7)
        canvas.drawString(30, height - 24, "RAY OS / АНАЛИТИКА СПРОСА")
        canvas.setFillColor(color(MUTED))
        canvas.setFont(regular, 6.5)
        canvas.drawRightString(
            width - 30,
            height - 24,
            f"{metadata['report_id']}  |  {metadata['model']}  |  данные {metadata['data_as_of']}",
        )
        canvas.drawString(30, 13, f"Предупреждений: {warning_count}")
        canvas.drawRightString(width - 30, 13, f"Страница {canvas.getPageNumber()}")
        canvas.setTitle(f"Отчёт Ray OS: {document['strategy']['title']}")
        canvas.setAuthor("Ray OS")
        canvas.restoreState()

    styles = {
        "hero": ParagraphStyle(
            "Hero",
            fontName=bold,
            fontSize=31,
            leading=34,
            textColor=color(WHITE),
            spaceAfter=13,
        ),
        "h1": ParagraphStyle(
            "H1",
            fontName=bold,
            fontSize=21,
            leading=24,
            textColor=color(WHITE),
            spaceAfter=10,
        ),
        "h2": ParagraphStyle(
            "H2",
            fontName=bold,
            fontSize=12,
            leading=15,
            textColor=color(LIME),
            spaceAfter=6,
        ),
        "body": ParagraphStyle(
            "Body",
            fontName=regular,
            fontSize=9,
            leading=13,
            textColor=color(POWDER),
            spaceAfter=6,
        ),
        "small": ParagraphStyle(
            "Small",
            fontName=regular,
            fontSize=7,
            leading=9,
            textColor=color(MUTED),
        ),
        "metric": ParagraphStyle(
            "Metric",
            fontName=bold,
            fontSize=18,
            leading=19,
            textColor=color(WHITE),
            alignment=1,
        ),
        "metric_label": ParagraphStyle(
            "MetricLabel",
            fontName=regular,
            fontSize=7,
            leading=9,
            textColor=color(MUTED),
            alignment=1,
        ),
        "cell": ParagraphStyle(
            "Cell",
            fontName=regular,
            fontSize=6.7,
            leading=8,
            textColor=color(WHITE),
        ),
        "cell_bold": ParagraphStyle(
            "CellBold",
            fontName=bold,
            fontSize=6.7,
            leading=8,
            textColor=color(WHITE),
        ),
    }

    def paragraph(value: Any, style: str = "body") -> Paragraph:
        return Paragraph(escape(str(value)), styles[style])

    def metric_row(items: Sequence[tuple[str, str, str]]) -> Table:
        cards = []
        for label, value, accent in items:
            card = Table(
                [[paragraph(value, "metric")], [paragraph(label.upper(), "metric_label")]],
                colWidths=[2.45 * inch],
                rowHeights=[0.44 * inch, 0.27 * inch],
            )
            card.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), color(CHARCOAL)),
                        ("BOX", (0, 0), (-1, -1), 1.1, color(accent)),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 8),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ]
                )
            )
            cards.append(card)
        outer = Table([cards], colWidths=[2.58 * inch] * len(cards), hAlign="LEFT")
        outer.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return outer

    header_style = TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), color(CORAL)),
            ("TEXTCOLOR", (0, 0), (-1, 0), color(NAVY)),
            ("FONTNAME", (0, 0), (-1, 0), bold),
            ("FONTSIZE", (0, 0), (-1, 0), 7),
            ("BACKGROUND", (0, 1), (-1, -1), color(CHARCOAL)),
            ("TEXTCOLOR", (0, 1), (-1, -1), color(WHITE)),
            ("FONTNAME", (0, 1), (-1, -1), regular),
            ("FONTSIZE", (0, 1), (-1, -1), 6.7),
            ("GRID", (0, 0), (-1, -1), 0.35, color(GRID)),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
    )

    portfolio = document["portfolio"]
    metrics = document["metrics"]
    story: list[Any] = []

    story.extend(
        [
            Spacer(1, 0.48 * inch),
            paragraph("RAY OS // СИГНАЛ ЗАПАСОВ НА 30 ДНЕЙ", "h2"),
            Paragraph(
                f"{escape(document['strategy']['title'])}<br/><font color='{CORAL}'>Спрос, превращённый в действие.</font>",
                styles["hero"],
            ),
            paragraph(document["strategy"]["description"], "body"),
            Spacer(1, 0.24 * inch),
            metric_row(
                [
                    ("Товаров в контуре", _format_number(portfolio["sku_count"]), CORAL),
                    ("Спрос q50 / 30 дней", _format_number(portfolio["forecast_30_q50"]), LIME),
                    ("К заказу, шт.", _format_number(portfolio["recommended_order_units"]), POWDER),
                    ("Ошибка WAPE", f"{metrics['wape_pct']:.1f}%", CORAL),
                ]
            ),
            Spacer(1, 0.28 * inch),
            paragraph(
                f"Период: {metadata['forecast_start']} - {metadata['forecast_end']}  /  "
                f"данные на {metadata['data_as_of']}  /  стратегия {document['report_type'].upper()}",
                "body",
            ),
            paragraph(
                "Правило решения: накопленный выбранный квантиль за срок поставки минус остаток и товар в пути. "
                "Неопределённость показывается отдельно и не учитывается дважды.",
                "small",
            ),
            PageBreak(),
        ]
    )

    story.extend(
        [
            paragraph("Пульс портфеля", "h1"),
            metric_row(
                [
                    ("Товаров под риском", _format_number(portfolio["at_risk_skus"]), CORAL),
                    ("Сумма заказа", _format_money(portfolio["recommended_order_value"]), LIME),
                    ("Стоимость запаса", _format_money(portfolio["on_hand_value"]), POWDER),
                    ("Покрытие q10-q90", f"{metrics['coverage_pct']:.1f}%", CORAL),
                ]
            ),
            Spacer(1, 0.12 * inch),
            Image(str(charts["portfolio"]), width=10.6 * inch, height=3.35 * inch),
            paragraph(
                "Коралловая линия — рабочий план q50; голубая область — эмпирический диапазон неопределённости q10-q90.",
                "small",
            ),
            PageBreak(),
        ]
    )

    story.extend([paragraph("Приоритетные действия", "h1"), paragraph("Включены все товары; позиции с наибольшим влиянием показаны первыми.", "body")])
    action_rows: list[list[Any]] = [
        ["ID", "ТОВАР", "КАТЕГОРИЯ", "ОСТАТОК", "В ПУТИ", "СРОК", "СПРОС LT", "БУФЕР", "ЗАКАЗ", "ДЕЙСТВИЕ"]
    ]
    for item in document["product_actions"]:
        action_rows.append(
            [
                str(item["product_id"]),
                paragraph(_short(item["product_name"], 25), "cell_bold"),
                paragraph(_short(item["category"], 18), "cell"),
                _format_number(item["on_hand"]),
                _format_number(item["in_transit"]),
                str(item["lead_time_days"]),
                _format_number(item["lead_time_demand"]),
                _format_number(item["safety_stock"]),
                _format_number(item["order_quantity"]),
                item["action"],
            ]
        )
    action_table = LongTable(
        action_rows,
        repeatRows=1,
        colWidths=[
            0.4 * inch,
            2.3 * inch,
            1.65 * inch,
            0.7 * inch,
            0.7 * inch,
            0.45 * inch,
            0.95 * inch,
            0.8 * inch,
            0.75 * inch,
            0.75 * inch,
        ],
        hAlign="LEFT",
    )
    action_table.setStyle(header_style)
    action_table.setStyle(
        TableStyle(
            [
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ("FONTSIZE", (0, 1), (-1, -1), 6.2),
            ]
        )
    )
    for row_index, item in enumerate(document["product_actions"], start=1):
        if item["order_quantity"] > 0:
            action_table.setStyle(
                TableStyle(
                    [
                        ("TEXTCOLOR", (8, row_index), (9, row_index), color(LIME)),
                        ("FONTNAME", (8, row_index), (9, row_index), bold),
                    ]
                )
            )
    story.extend(
        [
            action_table,
            Spacer(1, 0.12 * inch),
            paragraph("БУФЕР — разница q90 и q50 за срок поставки. Он показан справочно и повторно к заказу не добавляется.", "small"),
            PageBreak(),
        ]
    )

    story.extend(
        [
            paragraph("Спрос по категориям", "h1"),
            Image(str(charts["category"]), width=7.7 * inch, height=3.45 * inch),
        ]
    )
    category_rows: list[list[Any]] = [["КАТЕГОРИЯ", "ТОВАРЫ", "Q50 / 30 ДН.", "Q90 / 30 ДН.", "ЗАКАЗ", "СУММА ЗАКАЗА", "ПОД РИСКОМ"]]
    for item in document["category_summaries"]:
        category_rows.append(
            [
                paragraph(item["category"], "cell_bold"),
                str(item["sku_count"]),
                _format_number(item["forecast_30_q50"]),
                _format_number(item["forecast_30_q90"]),
                _format_number(item["order_quantity"]),
                _format_money(item["order_value"]),
                str(item["at_risk_skus"]),
            ]
        )
    category_table = LongTable(
        category_rows,
        repeatRows=1,
        colWidths=[2.15 * inch, 0.65 * inch, 1.25 * inch, 1.25 * inch, 0.9 * inch, 1.55 * inch, 0.75 * inch],
        hAlign="LEFT",
    )
    category_table.setStyle(header_style)
    story.extend([Spacer(1, 0.1 * inch), category_table, PageBreak()])

    quality = document["quality"]
    story.extend(
        [
            paragraph("Методология и качество сигнала", "h1"),
            metric_row(
                [
                    ("Качество данных", f"{quality['score']:.0f}/100", LIME),
                    ("Ошибка MAE", f"{metrics['mae']:.1f}", POWDER),
                    ("Смещение", f"{metrics['bias_pct']:+.1f}%", CORAL),
                    ("Возраст данных", f"{metadata['staleness_days']} дн.", CORAL),
                ]
            ),
            Spacer(1, 0.2 * inch),
            paragraph("Как формируется сигнал", "h2"),
            paragraph("1. Спрос при отсутствии товара восстанавливается только внутри того же product_id.", "body"),
            paragraph("2. Устойчивый недавний спрос и день недели формируют q50.", "body"),
            paragraph("3. Фактические промо и праздники задают ограниченный uplift; остатки ошибок формируют q10-q90.", "body"),
            paragraph("4. Исторический спрос ограничивает все квантили; прогнозный остаток не бывает отрицательным.", "body"),
            paragraph("5. Holdout разделён по целым датам, поэтому одна дата не попадает одновременно в train и test.", "body"),
            Spacer(1, 0.12 * inch),
            paragraph("Ограничения и предупреждения", "h2"),
        ]
    )
    warning_rows = [["#", "ПРЕДУПРЕЖДЕНИЕ"]]
    if document["warnings"]:
        warning_rows.extend(
            [[str(index), paragraph(warning, "cell")] for index, warning in enumerate(document["warnings"], start=1)]
        )
    else:
        warning_rows.append(["-", paragraph("Существенных предупреждений нет.", "cell")])
    warning_table = LongTable(
        warning_rows,
        repeatRows=1,
        colWidths=[0.35 * inch, 9.9 * inch],
        hAlign="LEFT",
    )
    warning_table.setStyle(header_style)
    story.append(warning_table)

    pdf = SimpleDocTemplate(
        str(destination),
        pagesize=page_size,
        leftMargin=30,
        rightMargin=30,
        topMargin=42,
        bottomMargin=31,
        title=f"Отчёт Ray OS: {document['strategy']['title']}",
        author="Ray OS",
    )
    pdf.build(story, onFirstPage=on_page, onLaterPages=on_page)


def _render_pptx(document: Mapping[str, Any], destination: Path, charts: Mapping[str, Path]) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    def rgb(value: str) -> RGBColor:
        value = value.lstrip("#")
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]
    metadata = document["metadata"]
    metrics = document["metrics"]
    portfolio = document["portfolio"]

    def add_text(
        slide,
        text: Any,
        x: float,
        y: float,
        width: float,
        height: float,
        *,
        size: float = 14,
        color: str = WHITE,
        bold: bool = False,
        align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP,
    ):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.02)
        frame.margin_top = frame.margin_bottom = Inches(0.02)
        frame.vertical_anchor = valign
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = str(text)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
        return shape

    def add_footer(slide, page_number: int) -> None:
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.45), Inches(7.12), Inches(12.43), Inches(0.012))
        line.fill.solid()
        line.fill.fore_color.rgb = rgb(GRID)
        line.line.fill.background()
        add_text(
            slide,
            f"{metadata['report_id']}  |  {metadata['model']}  |  данные {metadata['data_as_of']}  |  предупреждения {len(document['warnings'])}",
            0.48,
            7.18,
            10.9,
            0.18,
            size=6.5,
            color=MUTED,
        )
        add_text(slide, f"{page_number}/6", 11.95, 7.18, 0.85, 0.18, size=6.5, color=MUTED, align=PP_ALIGN.RIGHT)

    def new_slide(title: str | None, kicker: str, page_number: int):
        slide = presentation.slides.add_slide(blank_layout)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(NAVY)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.48), Inches(0.42), Inches(0.54), Inches(0.08))
        accent.fill.solid()
        accent.fill.fore_color.rgb = rgb(CORAL)
        accent.line.fill.background()
        add_text(slide, kicker.upper(), 1.12, 0.32, 4.8, 0.3, size=8, color=POWDER, bold=True)
        if title:
            add_text(slide, title, 0.48, 0.75, 12.2, 0.55, size=25, color=WHITE, bold=True)
        add_footer(slide, page_number)
        return slide

    def add_card(slide, label: str, value: str, x: float, y: float, width: float, accent: str) -> None:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(width), Inches(1.15))
        card.fill.solid()
        card.fill.fore_color.rgb = rgb(CHARCOAL)
        card.line.color.rgb = rgb(accent)
        card.line.width = Pt(1.2)
        add_text(slide, value, x + 0.17, y + 0.17, width - 0.34, 0.43, size=22, bold=True)
        add_text(slide, label.upper(), x + 0.17, y + 0.76, width - 0.34, 0.2, size=7, color=MUTED, bold=True)

    def style_table(table, widths: Sequence[float], header_color: str = CORAL) -> None:
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for row_index, row in enumerate(table.rows):
            row.height = Inches(0.34 if row_index == 0 else 0.31)
            for cell in row.cells:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(header_color if row_index == 0 else CHARCOAL)
                cell.margin_left = cell.margin_right = Inches(0.06)
                cell.margin_top = cell.margin_bottom = Inches(0.02)
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(7 if row_index == 0 else 7.5)
                        run.font.bold = row_index == 0
                        run.font.color.rgb = rgb(NAVY if row_index == 0 else WHITE)

    slide = new_slide(None, "Ray OS / Аналитика спроса", 1)
    add_text(slide, document["strategy"]["title"], 0.65, 1.28, 11.9, 0.82, size=38, bold=True)
    add_text(slide, "Спрос, превращённый в действие.", 0.65, 2.18, 11.2, 0.48, size=21, color=CORAL, bold=True)
    add_text(slide, document["strategy"]["description"], 0.67, 2.9, 8.8, 0.7, size=14, color=POWDER)
    add_card(slide, "Товаров в контуре", _format_number(portfolio["sku_count"]), 0.65, 4.22, 2.75, CORAL)
    add_card(slide, "Спрос q50 / 30 дней", _format_number(portfolio["forecast_30_q50"]), 3.58, 4.22, 2.75, LIME)
    add_card(slide, "К заказу, шт.", _format_number(portfolio["recommended_order_units"]), 6.51, 4.22, 2.75, POWDER)
    add_card(slide, "Ошибка WAPE", f"{metrics['wape_pct']:.1f}%", 9.44, 4.22, 2.75, CORAL)
    add_text(
        slide,
        f"{metadata['forecast_start']} - {metadata['forecast_end']}   /   {document['report_type'].upper()}",
        0.67,
        5.78,
        11.5,
        0.3,
        size=10,
        color=MUTED,
    )

    slide = new_slide("Пульс портфеля", "Рабочий сигнал", 2)
    add_card(slide, "Товаров под риском", _format_number(portfolio["at_risk_skus"]), 0.5, 1.42, 2.9, CORAL)
    add_card(slide, "Сумма заказа", _format_money(portfolio["recommended_order_value"]), 3.62, 1.42, 2.9, LIME)
    add_card(slide, "Стоимость запаса", _format_money(portfolio["on_hand_value"]), 6.74, 1.42, 2.9, POWDER)
    add_card(slide, "Покрытие q10-q90", f"{metrics['coverage_pct']:.1f}%", 9.86, 1.42, 2.9, CORAL)
    slide.shapes.add_picture(str(charts["portfolio"]), Inches(0.5), Inches(2.82), width=Inches(12.25), height=Inches(3.95))

    slide = new_slide("Спрос по категориям", "Где сосредоточен объём", 3)
    slide.shapes.add_picture(str(charts["category"]), Inches(0.5), Inches(1.45), width=Inches(8.05), height=Inches(4.95))
    y_position = 1.48
    for item, accent in zip(document["category_summaries"], [CORAL, LIME, POWDER, "#D7B5FF"]):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.82), Inches(y_position), Inches(3.95), Inches(1.02))
        card.fill.solid()
        card.fill.fore_color.rgb = rgb(CHARCOAL)
        card.line.color.rgb = rgb(accent)
        add_text(slide, _short(item["category"], 30), 9.02, y_position + 0.12, 3.45, 0.22, size=9, color=MUTED, bold=True)
        add_text(slide, _format_number(item["forecast_30_q50"]), 9.02, y_position + 0.42, 1.5, 0.35, size=18, bold=True)
        add_text(slide, f"заказ {item['order_quantity']}", 10.72, y_position + 0.52, 1.5, 0.22, size=8, color=accent, bold=True)
        y_position += 1.18

    slide = new_slide("Приоритетные действия", "Что делать дальше", 4)
    top_actions = document["product_actions"][:10]
    table_shape = slide.shapes.add_table(len(top_actions) + 1, 7, Inches(0.5), Inches(1.48), Inches(12.25), Inches(4.9))
    table = table_shape.table
    headers = ["ТОВАР", "КАТЕГОРИЯ", "ОСТАТОК", "В ПУТИ", "СПРОС LT", "ЗАКАЗ", "ДЕЙСТВИЕ"]
    for column, header in enumerate(headers):
        table.cell(0, column).text = header
    for row, item in enumerate(top_actions, start=1):
        values = [
            f"{item['product_id']}  {_short(item['product_name'], 22)}",
            _short(item["category"], 16),
            _format_number(item["on_hand"]),
            _format_number(item["in_transit"]),
            _format_number(item["lead_time_demand"]),
            _format_number(item["order_quantity"]),
            item["action"],
        ]
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    style_table(table, [2.7, 2.0, 1.05, 1.05, 1.45, 1.05, 1.25])
    add_text(
        slide,
        f"Квантиль решения: {document['strategy']['quantile']}  |  страховой буфер виден, но не учитывается дважды.",
        0.52,
        6.56,
        11.8,
        0.28,
        size=8,
        color=MUTED,
    )

    slide = new_slide("Портфель", "20 товаров, единая карта решений", 5)
    slide.shapes.add_picture(str(charts["scatter"]), Inches(0.5), Inches(1.43), width=Inches(8.2), height=Inches(5.12))
    add_card(slide, "Товаров в контуре", _format_number(portfolio["sku_count"]), 9.0, 1.5, 3.75, LIME)
    add_card(slide, "Позиций к заказу", _format_number(portfolio["order_skus"]), 9.0, 2.9, 3.75, CORAL)
    add_card(slide, "Спрос q90", _format_number(portfolio["forecast_30_q90"]), 9.0, 4.3, 3.75, POWDER)
    add_text(slide, "Размер круга = рекомендуемый заказ. Коралловый = требуется действие.", 9.04, 5.78, 3.55, 0.46, size=9, color=MUTED)

    slide = new_slide("Методология и ограничения", "Как читать сигнал", 6)
    quality = document["quality"]
    add_card(slide, "Качество данных", f"{quality['score']:.0f}/100", 0.5, 1.48, 2.85, LIME)
    add_card(slide, "Ошибка MAE", f"{metrics['mae']:.1f}", 3.53, 1.48, 2.85, POWDER)
    add_card(slide, "Смещение", f"{metrics['bias_pct']:+.1f}%", 6.56, 1.48, 2.85, CORAL)
    add_card(slide, "Возраст данных", f"{metadata['staleness_days']} дн.", 9.59, 1.48, 2.85, CORAL)
    methodology = (
        "01  OOS-спрос восстанавливается внутри одного товара\n"
        "02  Свежие данные и день недели формируют q50\n"
        "03  Фактические события задают ограниченный эффект\n"
        "04  Остатки ошибок формируют диапазон q10-q90\n"
        "05  Заказ = квантиль за срок - остаток - товар в пути"
    )
    add_text(slide, methodology, 0.55, 3.05, 5.85, 2.55, size=12, color=POWDER, bold=True)
    warnings = document["warnings"] or ["Существенных предупреждений нет."]
    warning_text = "ПРЕДУПРЕЖДЕНИЯ\n" + "\n".join(f"- {_short(item, 95)}" for item in warnings[:7])
    if len(warnings) > 7:
        warning_text += f"\n- ещё {len(warnings) - 7} в PDF"
    add_text(slide, warning_text, 6.75, 3.04, 5.55, 2.95, size=9.5, color=CORAL, bold=False)
    add_text(
        slide,
        "Последние уникальные даты целиком вынесены в holdout. Покрытие измеряется по эмпирическому диапазону q10-q90.",
        0.55,
        6.25,
        11.8,
        0.4,
        size=8,
        color=MUTED,
    )
    presentation.save(str(destination))


def generate_report_bundle(
    report_type: str = "standard",
    formats: Sequence[str] | str | None = None,
) -> dict[str, Any]:
    """Build one report document and render every requested artifact from it."""
    normalized_formats = _normalize_formats(formats)
    _check_dependencies(normalized_formats)
    document = build_report_document(report_type)
    output_dir = _output_directory()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    short_id = re.sub(r"[^A-Za-z0-9]", "", document["report_id"])[-6:].lower()
    safe_type = re.sub(r"[^a-z0-9_-]", "", str(report_type).lower()) or "standard"
    base_name = f"ray_os_{safe_type}_{timestamp}_{short_id}"
    artifacts: list[dict[str, str]] = []
    created_paths: list[Path] = []

    try:
        with tempfile.TemporaryDirectory(prefix=f"ray-os-{short_id}-") as temp_name:
            charts = _create_charts(document, Path(temp_name))
            for report_format in normalized_formats:
                destination = output_dir / f"{base_name}.{report_format}"
                created_paths.append(destination)
                if report_format == "pdf":
                    _render_pdf(document, destination, charts)
                else:
                    _render_pptx(document, destination, charts)
                artifacts.append(
                    {
                        "format": report_format,
                        "filename": destination.name,
                        "path": str(destination),
                    }
                )
    except Exception:
        for path in created_paths:
            path.unlink(missing_ok=True)
        raise

    return {
        "report_id": document["report_id"],
        "artifacts": artifacts,
        "metadata": document["metadata"],
        "warnings": document["warnings"],
    }
